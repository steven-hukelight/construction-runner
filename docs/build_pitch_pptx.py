#!/usr/bin/env python3
"""Build Construction Runner Future Upgrades pitch deck."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "user-guide" / "screenshots"
OUT = ROOT / "Future-Upgrades-Pitch.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
LIGHT_BLUE = RGBColor(0x93, 0xC5, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
SOFT = RGBColor(0xF8, 0xFA, 0xFC)
TIP = RGBColor(0xEF, 0xF6, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    return s


def chrome(slide):
    rect(slide, 0, 0, W, H, WHITE)
    rect(slide, 0, 0, W, Inches(0.12), BLUE)
    foot = rect(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    p = foot.text_frame.paragraphs[0]
    p.text = "Construction Runner  ·  Future upgrades  ·  www.construction-runner.com"
    set_run(p.runs[0], 10, False, WHITE)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide)
    return slide


def title_box(slide, text, kicker=None):
    top = 0.35
    if kicker:
        k = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(0.4))
        p = k.text_frame.paragraphs[0]
        p.text = kicker
        set_run(p.runs[0], 14, True, BLUE)
        top += 0.4
    box = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = text
    set_run(p.runs[0], 30, True, NAVY)
    return top + 0.85


def bullets(slide, items, top, left=0.55, width=12.3, size=19):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        bullet = p.add_run()
        bullet.text = "•  "
        set_run(bullet, size, True, BLUE)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run()
            r.text = lead
            set_run(r, size, True, INK)
            r = p.add_run()
            r.text = rest
            set_run(r, size, False, INK)
        else:
            r = p.add_run()
            r.text = item
            set_run(r, size, False, INK)


def callout(slide, text, top, label="Why it matters"):
    box = rect(slide, Inches(0.55), Inches(top), Inches(12.2), Inches(1.15), TIP,
               line=RGBColor(0xBF, 0xDB, 0xFE), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box.adjustments[0] = 0.12
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = label.upper()
    set_run(r, 11, True, BLUE)
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = text
    set_run(r, 17, False, INK)


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, Inches(6.9), W, Inches(0.6), BLUE)
    if (SHOTS / "icon.png").exists():
        slide.shapes.add_picture(str(SHOTS / "icon.png"), Inches(0.7), Inches(1.2), width=Inches(1.1))
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.6), Inches(12), Inches(2.4))
    p = box.text_frame.paragraphs[0]
    p.text = "Future upgrades"
    set_run(p.runs[0], 44, True, WHITE)
    p2 = box.text_frame.add_paragraph()
    p2.text = "From digital paperwork to a site that speaks up"
    set_run(p2.runs[0], 22, False, LIGHT_BLUE)
    sub = slide.shapes.add_textbox(Inches(0.7), Inches(5.2), Inches(11), Inches(1.2))
    p = sub.text_frame.paragraphs[0]
    p.text = "Construction Runner  ·  September 2026"
    set_run(p.runs[0], 16, False, WHITE)


def today(prs):
    s = new_slide(prs)
    top = title_box(s, "Where we are today")
    bullets(s, [
        "Phone sign-in that knows when someone is on site — and signs them out when they leave",
        "Inductions, RAMS and briefing sign-off, near-miss reporting",
        "Tasks, deliveries, assets, messaging and subcontractor onboarding",
        "A web dashboard for the office",
    ], top)
    callout(s, "That puts us level with the main UK competitors. The upgrades are about moving ahead of them.",
            5.5, label="The point")


def big_idea(prs):
    s = new_slide(prs)
    top = title_box(s, "We already know three things paperwork apps don't", kicker="The big idea")
    cards = [
        ("Where people are", "Live site sign-in and sign-out"),
        ("What they're qualified for", "Cards and certificates with expiry dates"),
        ("What they've read", "Inductions, RAMS and briefings"),
    ]
    card_w, gap = 3.9, 0.25
    for i, (head, body) in enumerate(cards):
        left = 0.55 + i * (card_w + gap)
        card = rect(s, Inches(left), Inches(top + 0.2), Inches(card_w), Inches(2.2), SOFT,
                    line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        card.adjustments[0] = 0.08
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = head
        set_run(r, 22, True, BLUE)
        p2 = tf.add_paragraph()
        p2.space_before = Pt(8)
        r = p2.add_run()
        r.text = body
        set_run(r, 17, False, INK)
    callout(s, "Every upgrade plugs into these three — so the app doesn't just store records, "
               "it notices when something is wrong and tells someone.", 5.3, label="So what")


def upgrade(prs, number, title, items, why=None):
    s = new_slide(prs)
    top = title_box(s, title, kicker=f"Upgrade {number}")
    bullets(s, items, top)
    if why:
        callout(s, why, 5.5)


def order(prs):
    s = new_slide(prs)
    top = title_box(s, "Suggested order")
    rows = [
        ("Step", "Upgrade", "Why this order"),
        ("1", "Offline sign-in fix", "Reliability of what customers already use"),
        ("2", "CSCS cards that stay checked", "Quick win; permits depend on it"),
        ("3", "Understanding checks on RAMS & briefings", "Our headline difference"),
        ("4", "Visitor QR sign-in", "High demand from buyers"),
        ("5", "Smart permits", "Builds on steps 2 and 3"),
    ]
    table = s.shapes.add_table(len(rows), 3, Inches(0.55), Inches(top + 0.1),
                               Inches(12.2), Inches(0.6 * len(rows))).table
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(5.6)
    table.columns[2].width = Inches(5.5)
    for r_i, row in enumerate(rows):
        for c_i, text in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r_i == 0 else (WHITE if r_i % 2 else SOFT)
            p = cell.text_frame.paragraphs[0]
            p.text = text
            set_run(p.runs[0], 17, r_i == 0 or c_i == 0, WHITE if r_i == 0 else INK)


def closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, Inches(6.9), W, Inches(0.6), BLUE)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(3.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "IN ONE LINE"
    set_run(p.runs[0], 14, True, LIGHT_BLUE)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    p2.text = "Construction Runner already runs the site."
    set_run(p2.runs[0], 34, True, WHITE)
    p3 = tf.add_paragraph()
    p3.space_before = Pt(8)
    p3.text = ("Next, it proves the right people are in the right place, qualified and briefed "
               "— and speaks up when they're not.")
    set_run(p3.runs[0], 26, False, LIGHT_BLUE)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    today(prs)
    big_idea(prs)
    upgrade(prs, 1, "Prove workers understood, not just signed", [
        "After reading a RAMS or toolbox talk, the worker answers 3 quick questions from that exact document",
        "Different questions each time",
        "A bad question is pulled automatically if too many people get it wrong",
    ], why="A signature proves someone tapped \"I agree\". This proves they understood — stronger evidence "
           "for clients, insurers and the HSE. No UK competitor combines this with live site sign-in.")
    upgrade(prs, 2, "Permits that know where, who and when", [
        ("Qualified only: ", "only people with valid qualifications can be added to a permit"),
        ("Issuer leaves, permit pauses: ", "suspended automatically if the issuer leaves site"),
        ("Site map: ", "permits drawn on a map, with warnings when two clash"),
        ("Fire watch: ", "countdown after hot works, with reminders and a closing photo"),
    ], why="Other permit apps are a form with signatures. Ours checks the people, the place and the time.")
    upgrade(prs, 3, "Visitors with no iPad and no blind spots", [
        ("QR poster at the gate: ", "visitors sign in on their own phone — no kiosk to buy"),
        ("Linked to a host: ", "if the host leaves, they're reminded their visitor is still on site"),
        ("Fire roll-call: ", "visitors listed alongside workers"),
        ("Deliveries: ", "drivers matched to their booked slot automatically"),
    ])
    upgrade(prs, 4, "CSCS cards that stay checked", [
        ("Scan once: ", "the app tracks expiry and reminds the worker before it lapses"),
        ("Right person, right job: ", "shows who is on site today without the qualifications their job needs"),
        ("Spot problems: ", "flags the same card on two people, or an expired card presented"),
    ], why="Other apps check a card once at the gate. Ours keeps checking it against the work being done.")
    upgrade(prs, 5, "Works with no signal", [
        "Finish offline sign-in and sign-out",
        "Workers on poor-signal sites are still recorded accurately",
        "Everything syncs automatically when they reconnect",
    ])
    order(prs)
    closing(prs)
    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
