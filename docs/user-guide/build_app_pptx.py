#!/usr/bin/env python3
"""Build Construction Runner App Guide PowerPoint (for operatives)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "Construction-Runner-App-Guide.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    s.fill.solid()
    s.fill.fore_color.rgb = BLUE
    s.line.fill.background()
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.22), W, Inches(0.28))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    p = foot.text_frame.paragraphs[0]
    p.text = "Construction Runner  ·  App guide for operatives  ·  info@construction-runner.com"
    set_run(p.runs[0], 10, False, WHITE)


def title_box(slide, text, top=0.35, size=30):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(0.65))
    p = box.text_frame.paragraphs[0]
    p.text = text
    set_run(p.runs[0], size, True, NAVY)


def bullets(slide, items, top=1.15, width=12.3, size=17):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(width), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.space_after = Pt(8)
        set_run(p.runs[0], size, False, INK)


def picture(slide, name, left, top, width):
    path = SHOTS / name
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    bar(slide)
    return slide


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), W, Inches(0.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    if (SHOTS / "icon.png").exists():
        slide.shapes.add_picture(str(SHOTS / "icon.png"), Inches(0.7), Inches(1.2), width=Inches(1.1))
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.6), Inches(12), Inches(2.4))
    p = box.text_frame.paragraphs[0]
    p.text = "Your Construction Runner app"
    set_run(p.runs[0], 40, True, WHITE)
    p2 = box.text_frame.add_paragraph()
    p2.text = "A simple guide for operatives and anyone who uses the phone on site"
    set_run(p2.runs[0], 18, False, RGBColor(0x93, 0xC5, 0xFD))
    sub = slide.shapes.add_textbox(Inches(0.7), Inches(5.2), Inches(11), Inches(1.2))
    p = sub.text_frame.paragraphs[0]
    p.text = "iPhone & Android  ·  Version 1.0  ·  September 2026"
    set_run(p.runs[0], 16, False, WHITE)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)

    s = new_slide(prs)
    title_box(s, "This guide is for you")
    bullets(s, [
        "You work on site — this is the app you use every day",
        "You do not log in on the website. Use Construction Runner on your phone",
        "Sign in when you arrive, complete induction, read RAMS and briefings, do your tasks, sign out when you leave",
        "Ask your supervisor if something is waiting for approval",
    ])

    s = new_slide(prs)
    title_box(s, "First time")
    bullets(s, [
        "1. Install Construction Runner from the App Store or Google Play",
        "2. Sign in, or Register with the company invite code — then wait for approval",
        "3. Set Location to Always (iPhone) or Allow all the time (Android)",
        "4. Profile → My Inductions → complete induction for your site",
        "5. Profile → My Info → add an emergency contact",
    ])

    s = new_slide(prs)
    title_box(s, "Every day on site")
    bullets(s, [
        "1. Home → Check In / Out → pick your site → stand inside the boundary → Sign In",
        "2. Open Safety — Sign & acknowledge any new briefings or RAMS",
        "3. Do your Tasks. Use Deliveries, Assets or Messages if you need them",
        "4. When you leave: tap Sign Out, or walk out of the site fence (auto sign-out)",
    ])

    s = new_slide(prs)
    title_box(s, "Sign in to the app")
    picture(s, "20-app-login.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Location must be Always")
    bullets(s, [
        "Used only to detect the site boundary — not to track you around town",
        "iPhone: Allow While Using, then Change to Always Allow",
        "Android: Permissions → Location → Allow all the time",
        "While Using or Allow Once cannot sign you out after you leave — Sign In stays locked",
        "If asked, tap Continue, then Open Settings",
    ], size=16)

    s = new_slide(prs)
    title_box(s, "Always location")
    picture(s, "21-app-location.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Home")
    picture(s, "11-mobile-home.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Check In when you arrive")
    bullets(s, [
        "Home → Check In / Out → Select site",
        "Wait for Inside Site Boundary — you cannot check in from outside the fence",
        "Complete Site induction if asked, then Complete induction",
        "Accept RAMS if asked (Accept & Sign In)",
        "Tap Sign In",
    ], size=16)

    s = new_slide(prs)
    title_box(s, "Check In / Out")
    picture(s, "12-mobile-checkin.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Sign out when you leave")
    bullets(s, [
        "Open Check In / Out and tap Sign Out",
        "Or walk out of the geofence — You left the site — signed out automatically",
        "No signal? Sign Out can wait and sync later",
        "Do not stay signed in overnight if you have left site",
        "Day off? Mark as absent → Lock this day",
    ])

    s = new_slide(prs)
    title_box(s, "Induction for each site")
    bullets(s, [
        "Profile → My Inductions",
        "Not Started → Start Induction",
        "Completed lasts 365 days — Expired means do it again",
        "Read site safety, acknowledge RAMS, confirm site rules, Complete induction",
        "If the site requires induction, you cannot finish Sign In until this is done",
    ], size=16)

    s = new_slide(prs)
    title_box(s, "My Inductions")
    picture(s, "13-inductions.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Safety tab")
    picture(s, "22-app-safety.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Sign & acknowledge")
    picture(s, "23-app-briefing.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Near miss")
    picture(s, "24-app-nearmiss.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Tasks and the other tabs")
    picture(s, "26-app-tasks.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Profile")
    picture(s, "25-app-profile.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "If it is not working")
    bullets(s, [
        "Cannot sign in after register — wait for your supervisor to approve you",
        "Sign In locked / needs Always — set Location to Always, then come back",
        "Outside Site Boundary — walk onto the site and refresh location",
        "Asked for induction — Profile → My Inductions → Start Induction",
        "Did not auto sign out — Location is still While Using; Sign Out manually this time",
        "Forgot password — Forgot Password? on the sign-in screen",
        "Help: your supervisor, or info@construction-runner.com",
    ], size=16)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
