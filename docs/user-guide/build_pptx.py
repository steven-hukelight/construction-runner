#!/usr/bin/env python3
"""Build Construction Runner User Guide PowerPoint (16:9)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "Construction-Runner-User-Guide.pptx"

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
SOFT = RGBColor(0xF1, 0xF5, 0xF9)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def bar(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    s.fill.solid()
    s.fill.fore_color.rgb = BLUE
    s.line.fill.background()
    foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.22), W, Inches(0.28))
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    tf = foot.text_frame
    tf.margin_left = Inches(0.4)
    p = tf.paragraphs[0]
    p.text = "Construction Runner  ·  User Guide  ·  www.construction-runner.com"
    set_run(p.runs[0], 10, False, WHITE)


def title_box(slide, text, top=0.35, size=32):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = text
    set_run(p.runs[0], size, True, NAVY)
    return box


def body(slide, text, top=1.1, height=1.2, size=16):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(12.4), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    set_run(p.runs[0], size, False, INK)
    return box


def bullets(slide, items, top=1.15, width=12.3, size=16):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(width), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        set_run(p.runs[0], size, False, INK)
    return box


def picture(slide, name, left, top, width):
    path = SHOTS / name
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    bar(slide)
    return slide


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), W, Inches(0.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    if (SHOTS / "icon.png").exists():
        slide.shapes.add_picture(str(SHOTS / "icon.png"), Inches(0.7), Inches(1.1), width=Inches(1.1))
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(11.5), Inches(2.2))
    p = box.text_frame.paragraphs[0]
    p.text = "Construction Runner"
    set_run(p.runs[0], 44, True, WHITE)
    p2 = box.text_frame.add_paragraph()
    p2.text = "User Guide"
    set_run(p2.runs[0], 32, False, RGBColor(0x93, 0xC5, 0xFD))
    sub = slide.shapes.add_textbox(Inches(0.7), Inches(5.0), Inches(11), Inches(1.2))
    p = sub.text_frame.paragraphs[0]
    p.text = "Web admin and mobile app  ·  Version 1.1  ·  September 2026"
    set_run(p.runs[0], 16, False, WHITE)
    p = sub.text_frame.add_paragraph()
    p.text = "www.construction-runner.com   ·   info@construction-runner.com"
    set_run(p.runs[0], 14, False, RGBColor(0xCB, 0xD5, 0xE1))


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    cover(prs)

    s = new_slide(prs)
    title_box(s, "What this guide covers")
    bullets(s, [
        "How people join, sign in and reset a password",
        "Roles: Super Admin, Site Admin, Supervisor, Operative, Subcontractor admin",
        "Sites as folders — documents and tasks stay on the site you pick",
        "Users, pending approvals, and ticking multiple sites for admins & supervisors",
        "Attendance on web and GPS check-in on the app",
        "Safety: RAMS, briefings, induction safety, site rules, alerts, near miss",
        "Missing Info, induction (365-day expiry), and the Construction Runner mobile app",
    ])

    s = new_slide(prs)
    title_box(s, "Two products, one company")
    bullets(s, [
        "Web admin — www.construction-runner.com → Member Login",
        "For Super Admins, Site Admins, Supervisors and Subcontractor admins",
        "Mobile app — Construction Runner on iPhone and Android",
        "Operatives use the app only (web login is not available for operatives)",
        "Same account: email, password, then admin approval before first sign-in",
    ], top=1.1, width=6.2)
    picture(s, "01-landing.png", 7.0, 1.15, 5.8)

    s = new_slide(prs)
    title_box(s, "Website and contact")
    picture(s, "05-contact.png", 0.5, 1.15, 12.3)

    s = new_slide(prs)
    title_box(s, "Roles at a glance")
    bullets(s, [
        "Super Admin — whole company, create sites, grant any role including Super Admin",
        "Site Admin — only ticked sites; cannot create sites or grant Super Admin",
        "Supervisor — same site ticks; live attendance, tasks, safety; app Supervisor Tools",
        "Operative — app: geofence sign-in, induction, RAMS, tasks, near miss",
        "Subcontractor admin — onboard their operatives on linked sites",
        "Site Admins and Supervisors can work on more than one site — tick every site they need",
    ])

    s = new_slide(prs)
    title_box(s, "Sign in to the web admin")
    picture(s, "02-login.png", 0.5, 1.05, 12.3)

    s = new_slide(prs)
    title_box(s, "Create an account")
    bullets(s, [
        "Join with invite code — code from Settings → Company",
        "Or Request new company — pending superuser approval",
        "After submit: “Account request submitted. Pending approval.”",
        "They cannot use web or app until someone approves them",
    ], top=1.1, width=6.0, size=15)
    picture(s, "03-register.png", 6.7, 1.1, 6.1)

    s = new_slide(prs)
    title_box(s, "Forgot password")
    bullets(s, [
        "Sign in → Forgot password? → enter email → Send reset link",
        "Open the email, set a new password, sign in",
        "Super Admins can also Users → Send reset email",
        "If the link has expired, request a new one",
    ], top=1.1, width=6.0, size=15)
    picture(s, "04-forgot-password.png", 6.7, 1.1, 6.1)

    s = new_slide(prs)
    title_box(s, "Sites — one folder per site")
    bullets(s, [
        "Super Admin: Add New Site — name, map, geofence radius, show on mobile map",
        "Tick Induction required if people must induct before Sign In",
        "Site Admin / Supervisor: assigned sites only — no Add Site",
        "Open a site for Details, Subcontractors and Induction tabs",
        "RAMS, briefings, tasks, alerts and rules stay on the site you selected",
    ], top=1.05, width=12.2, size=16)

    s = new_slide(prs)
    title_box(s, "Sites")
    picture(s, "06-sites.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Pending approvals")
    bullets(s, [
        "Until approved, the person cannot sign in on web or app",
        "Assign role: Operative, Supervisor, Site Admin, Super Admin",
        "Only a Super Admin can grant Super Admin",
        "Supervisor and Site Admin: tick all sites they can access — at least one",
        "Approve or Reject",
    ], top=1.05, width=12.2, size=16)

    s = new_slide(prs)
    title_box(s, "Pending approvals")
    picture(s, "08-approvals.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Assign sites — admins and supervisors")
    bullets(s, [
        "Users → row menu → Assign sites…",
        "Same picker for Site Admin and Supervisor",
        "Tick every site they work on — they may cross over sites",
        "Select all / Clear, then Save",
        "Invite User shows the same ticks when the role needs sites",
    ], top=1.05, width=12.2, size=16)

    s = new_slide(prs)
    title_box(s, "Assign sites")
    picture(s, "07-assign-sites.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Attendance (web)")
    bullets(s, [
        "Live — who is signed in; archives at midnight UK",
        "Role Call — fire roll; Signed in / Signed out; Export CSV or PDF",
        "Record attendance — optional manual web sign-in/out",
        "GPS check-in on the app is the source of truth on site",
        "Auto sign-out can show as geofence exit or left site",
    ], top=1.05, width=12.2, size=16)

    s = new_slide(prs)
    title_box(s, "Attendance")
    picture(s, "09-attendance.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Day-to-day work")
    bullets(s, [
        "Tasks — Add Task: title, people, required Site, due date",
        "Deliveries — reference, supplier, site, scheduled time; photos on the app",
        "Assets — equipment, assign, inspections",
        "Messages — only threads you are in; New → pick people (optional company broadcast)",
        "Always choose the correct site before you save",
    ])

    s = new_slide(prs)
    title_box(s, "Safety")
    bullets(s, [
        "RAMS — upload PDF for one site; operatives acknowledge on the app",
        "Briefings — toolbox talks; Sign & acknowledge; download reports",
        "Induction safety — company pack shown during induction (not Site Rules)",
        "COSHH — substances, symbols, PPE",
        "Site Rules — PPE, emergency procedures, conduct — per site",
        "Alerts — Info / Warning / Critical + site",
        "Near Miss — from the app; Mark reviewed; badge on Safety until cleared",
    ])

    s = new_slide(prs)
    title_box(s, "RAMS")
    picture(s, "10-rams.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Induction and Missing Info")
    bullets(s, [
        "Per site. Expires 365 days after completion",
        "Missing Info — people missing emergency contact or medical details",
        "Induction steps: Induction safety pack → RAMS → site rules → Complete",
        "Site → Induction tab: mark complete or reset",
        "App: Profile → My Inductions → Start Induction",
        "If induction is required, Check In will not finish until it is completed",
    ])

    s = new_slide(prs)
    title_box(s, "My Inductions")
    picture(s, "13-inductions.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Mobile app — Home")
    picture(s, "11-mobile-home.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "Check In / Out")
    bullets(s, [
        "Home → Check In / Out → Select site",
        "Wait for Inside Site Boundary, then Sign In",
        "Complete induction and Accept & Sign In if asked",
        "Sign Out, or leave the geofence for automatic sign-out",
        "iPhone: Location = Always    Android: Allow all the time",
        "While Using / Allow Once will not sign you out after you leave",
    ], size=16)

    s = new_slide(prs)
    title_box(s, "Check In / Out")
    picture(s, "12-mobile-checkin.png", 0.4, 1.05, 12.5)

    s = new_slide(prs)
    title_box(s, "App tabs")
    bullets(s, [
        "Home — Check In / Out, tasks, briefings",
        "Safety — briefings, RAMS, site rules, alerts, near miss, COSHH",
        "Deliveries — submit delivery, proof photos",
        "Assets — My Assets, inspections",
        "Messages and Tasks",
        "Menu: Profile, My Inductions, Settings, Sign Out",
        "Sign In needs a live connection and GPS — there is no web Offline queue",
    ])

    s = new_slide(prs)
    title_box(s, "Subcontractors")
    bullets(s, [
        "Main contractor: Subcontractors → Invite Subcontractor → pick site → Generate invite code",
        "Partner: www.construction-runner.com/join — Join as subcontractor",
        "After approval they use Operative Onboarding for their own people",
        "This is separate from the company invite code used by your own staff",
    ])

    s = new_slide(prs)
    title_box(s, "If something goes wrong")
    bullets(s, [
        "Cannot sign in after registering — still pending approval",
        "Supervisor / Site Admin sees nothing — Assign sites and tick at least one",
        "Cannot Sign In on the app — Always location, inside geofence, induction, RAMS",
        "Did not auto sign out — Location is not set to Always",
        "Reset link expired — request a new one",
        "Wrong documents — check which site was selected on upload",
    ])

    s = new_slide(prs)
    title_box(s, "Support")
    bullets(s, [
        "Email: info@construction-runner.com",
        "Website: www.construction-runner.com",
        "Contact / demo: www.construction-runner.com/contact",
        "Privacy: www.construction-runner.com/legal/privacy-and-security",
        "This deck matches product labels as of September 2026 (guide v1.1, app 1.0.2)",
        "Public website figures are live screenshots; admin figures match current table chrome and status chips",
    ])

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
